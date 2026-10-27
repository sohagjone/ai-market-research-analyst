"""Export a Markdown market report to Word, Excel, PowerPoint and PDF."""
import io
import math
import re
from datetime import date

PALETTE = ["4F46E5", "7C3AED", "EC4899", "F59E0B", "14B8A6", "3B82F6", "EF4444", "10B981"]
DARK = "1F2340"
LIGHT_ROW = "F3F4FF"
BRAND = "AI Market Research Analyst"
DISCLAIMER = "AI-generated research. Verify key figures with primary sources before making decisions."


# ------------------------------------------------------------------ parsing
def clean(t):
    t = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1", t)
    t = t.replace("`", "")
    t = re.sub(r"(?<![\w*])\*(?!\*)(.+?)(?<!\*)\*(?![\w*])", r"\1", t)
    t = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"\1", t)
    return t.strip()


def split_bold(t):
    out = []
    for part in re.split(r"(\*\*.+?\*\*)", clean(t)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], True))
        else:
            out.append((part, False))
    return out or [("", False)]


def plain(t):
    return "".join(s for s, _ in split_bold(t)).replace("**", "")


def parse_markdown(md):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        if not s:
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            blocks.append({"t": "h", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                is_sep = all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) and any(cells)
                if not is_sep:
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                blocks.append({"t": "table", "rows": rows})
            continue
        m = re.match(r"^(\s*)[-*+•]\s+(.*)", raw)
        if m:
            blocks.append({"t": "bullet", "level": 1 if len(m.group(1)) >= 2 else 0, "text": m.group(2)})
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)", raw)
        if m:
            blocks.append({"t": "num", "n": m.group(2), "level": 1 if len(m.group(1)) >= 3 else 0,
                           "text": m.group(3)})
            i += 1
            continue
        blocks.append({"t": "p", "text": re.sub(r"^>\s*", "", s)})
        i += 1
    return blocks


def split_sections(blocks):
    """Returns (sections, title_from_h1). Sections split on H1/H2 headings."""
    title = None
    if blocks and blocks[0]["t"] == "h" and blocks[0]["level"] == 1:
        title = plain(blocks[0]["text"])
        blocks = blocks[1:]
    sections, cur = [], {"title": "Overview", "blocks": []}
    for b in blocks:
        if b["t"] == "h" and b["level"] <= 2:
            if cur["blocks"]:
                sections.append(cur)
            cur = {"title": plain(b["text"]), "blocks": []}
        else:
            cur["blocks"].append(b)
    if cur["blocks"]:
        sections.append(cur)
    return sections, title


def _num(v):
    v2 = v.replace(",", "").strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", v2):
        return float(v2) if "." in v2 else int(v2)
    return v


# ------------------------------------------------------------------ DOCX
def to_docx(md, title, subtitle):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    sections, h1 = split_sections(parse_markdown(md))
    title = title or h1 or "Market Research Report"

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def border_bottom(p, color):
        pPr = p._p.get_or_add_pPr()
        pb = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "10")
        b.set(qn("w:space"), "1")
        b.set(qn("w:color"), color)
        pb.append(b)
        pPr.append(pb)

    def runs(p, text, size=None, color=None, force_bold=False):
        for seg, bold in split_bold(text):
            r = p.add_run(seg)
            r.bold = bold or force_bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    # Cover banner
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.rows[0].cells[0]
    shade(cell, PALETTE[0])
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(255, 255, 255)
    for text, size in ((subtitle, 13), (f"Prepared {date.today():%B %d, %Y}  |  {BRAND}", 10)):
        if text:
            q = cell.add_paragraph()
            rr = q.add_run(text)
            rr.font.size = Pt(size)
            rr.font.color.rgb = RGBColor(0xE0, 0xE7, 0xFF)
    cell.add_paragraph().paragraph_format.space_after = Pt(6)
    strip = doc.add_table(rows=1, cols=4)
    for idx, c in enumerate(strip.rows[0].cells):
        shade(c, PALETTE[(idx + 1) % len(PALETTE)])
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        c.paragraphs[0].add_run(" ").font.size = Pt(3)
    doc.add_paragraph()

    for si, section in enumerate(sections):
        accent = PALETTE[si % len(PALETTE)]
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.keep_with_next = True
        runs(h, section["title"], size=16, color=accent, force_bold=True)
        border_bottom(h, accent)
        for b in section["blocks"]:
            t = b["t"]
            if t == "h":
                hp = doc.add_paragraph()
                hp.paragraph_format.space_before = Pt(8)
                hp.paragraph_format.keep_with_next = True
                runs(hp, b["text"], size=12.5, color=accent, force_bold=True)
            elif t == "p":
                runs(doc.add_paragraph(), b["text"])
            elif t == "bullet":
                runs(doc.add_paragraph(style="List Bullet 2" if b["level"] else "List Bullet"), b["text"])
            elif t == "num":
                np_ = doc.add_paragraph()
                np_.paragraph_format.left_indent = Inches(0.35 + 0.25 * b["level"])
                np_.paragraph_format.first_line_indent = Inches(-0.25)
                r = np_.add_run(f"{b['n']}.  ")
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(accent)
                runs(np_, b["text"])
            elif t == "table":
                rows = b["rows"]
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        c = table.rows[ri].cells[ci]
                        c.text = ""
                        para = c.paragraphs[0]
                        if ri == 0:
                            shade(c, accent)
                            runs(para, val, size=9.5, color="FFFFFF", force_bold=True)
                        else:
                            if ri % 2 == 0:
                                shade(c, LIGHT_ROW)
                            runs(para, val, size=9.5)
                doc.add_paragraph()

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(f"{BRAND}  |  {DISCLAIMER}")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ XLSX
def to_xlsx(md, title, subtitle):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sections, h1 = split_sections(parse_markdown(md))
    title = title or h1 or "Market Research Report"
    wb = Workbook()
    thin = Side(style="thin", color="D9DCF5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda c: PatternFill("solid", fgColor=c)
    white_bold = Font(bold=True, color="FFFFFF", name="Calibri")

    # --- Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = fill(PALETTE[0])
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 42
    ws.merge_cells("A2:D2")
    ws["A2"] = f"{subtitle}   |   Prepared {date.today():%B %d, %Y}"
    ws["A2"].font = Font(size=11, color="FFFFFF")
    ws["A2"].fill = fill(PALETTE[1])
    ws["A2"].alignment = Alignment(indent=1)
    ws.row_dimensions[2].height = 24
    row = 4
    ws.cell(row=row, column=1, value="Key points").font = Font(bold=True, size=13, color=PALETTE[2])
    row += 1
    exec_sec = next((s for s in sections if "executive" in s["title"].lower()), sections[0] if sections else None)
    if exec_sec:
        for b in exec_sec["blocks"]:
            if b["t"] in ("bullet", "num", "p"):
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                c = ws.cell(row=row, column=1, value="•  " + plain(b["text"]))
                c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
                ws.row_dimensions[row].height = max(18, 15 * math.ceil(len(plain(b["text"])) / 110))
                row += 1
    row += 1
    ws.cell(row=row, column=1, value="Sheets in this workbook").font = Font(bold=True, size=13, color=PALETTE[2])
    row += 1
    for col, w in zip("ABCD", (34, 30, 30, 30)):
        ws.column_dimensions[col].width = w
    index_row = row

    # --- Full report sheet
    rp = wb.create_sheet("Report")
    rp.sheet_view.showGridLines = False
    for ci, h in enumerate(("Section", "Type", "Content"), 1):
        c = rp.cell(row=1, column=ci, value=h)
        c.font, c.fill = white_bold, fill(PALETTE[0])
        c.alignment = Alignment(vertical="center", indent=1)
    rp.column_dimensions["A"].width = 32
    rp.column_dimensions["B"].width = 12
    rp.column_dimensions["C"].width = 110
    rp.freeze_panes = "A2"
    r = 2
    used = {"Summary", "Report"}
    table_sheets = []

    def unique_name(base):
        base = re.sub(r"[\[\]:*?/\\]", "", base)[:26].strip() or "Table"
        name, n = base, 2
        while name in used:
            name = f"{base[:23]} {n}"
            n += 1
        used.add(name)
        return name

    for si, s in enumerate(sections):
        accent = PALETTE[si % len(PALETTE)]
        for ci in range(1, 4):
            c = rp.cell(row=r, column=ci, value=s["title"] if ci == 1 else None)
            c.fill = fill(accent)
            c.font = white_bold
        r += 1
        for b in s["blocks"]:
            if b["t"] == "table":
                name = unique_name(s["title"])
                table_sheets.append(name)
                tws = wb.create_sheet(name)
                for ri, rowvals in enumerate(b["rows"], 1):
                    for ci, v in enumerate(rowvals, 1):
                        c = tws.cell(row=ri, column=ci, value=plain(v) if ri == 1 else _num(plain(v)))
                        c.border = border
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                        if ri == 1:
                            c.font, c.fill = white_bold, fill(accent)
                        elif ri % 2 == 0:
                            c.fill = fill(LIGHT_ROW)
                for ci in range(1, len(b["rows"][0]) + 1):
                    longest = max(len(str(tws.cell(row=ri, column=ci).value or "")) for ri in range(1, len(b["rows"]) + 1))
                    tws.column_dimensions[get_column_letter(ci)].width = min(max(14, longest + 3), 55)
                tws.freeze_panes = "A2"
                tws.sheet_properties.tabColor = accent
                rp.cell(row=r, column=2, value="Table").font = Font(italic=True, color="6B7280")
                rp.cell(row=r, column=3, value=f"See sheet '{name}'").font = Font(italic=True, color=accent)
                r += 1
            else:
                kind = {"h": "Heading", "p": "Text", "bullet": "Bullet", "num": "Item"}[b["t"]]
                rp.cell(row=r, column=2, value=kind).font = Font(color="6B7280")
                c = rp.cell(row=r, column=3, value=plain(b["text"]))
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if b["t"] == "h":
                    c.font = Font(bold=True, color=accent)
                rp.row_dimensions[r].height = max(16, 15 * math.ceil(len(plain(b["text"])) / 115))
                r += 1
    rp.sheet_properties.tabColor = PALETTE[0]
    ws.sheet_properties.tabColor = PALETTE[2]

    ws.cell(row=index_row, column=1, value="Summary").font = Font(bold=True)
    ws.cell(row=index_row, column=2, value="Title and key points")
    ws.cell(row=index_row + 1, column=1, value="Report").font = Font(bold=True)
    ws.cell(row=index_row + 1, column=2, value="Full report, section by section")
    for k, name in enumerate(table_sheets):
        ws.cell(row=index_row + 2 + k, column=1, value=name).font = Font(bold=True)
        ws.cell(row=index_row + 2 + k, column=2, value="Data table")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PPTX
def to_pptx(md, title, subtitle):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    sections, h1 = split_sections(parse_markdown(md))
    title = title or h1 or "Market Research Report"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    rgb = RGBColor.from_string

    def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        s.line.fill.background()
        return s

    def text(slide, x, y, w, h, value, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = value
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        r.font.color.rgb = rgb(color)
        return tb

    def rich(p, value, size, color=DARK):
        for seg, bold in split_bold(value):
            r = p.add_run()
            r.text = seg
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
            r.font.color.rgb = rgb(color)

    # --- Title slide
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, PALETTE[0])
    rect(s, 8.8, -1.8, 6.2, 6.2, PALETTE[1], MSO_SHAPE.OVAL)
    rect(s, 10.4, 4.3, 4.6, 4.6, PALETTE[2], MSO_SHAPE.OVAL)
    rect(s, 5.5, 6.3, 3.0, 3.0, PALETTE[3], MSO_SHAPE.OVAL)
    rect(s, 0.9, 2.2, 0.12, 2.4, PALETTE[3])
    text(s, 1.2, 1.2, 10, 0.5, "MARKET RESEARCH REPORT", 14, True, "FDE68A")
    text(s, 1.2, 2.1, 9.5, 2.0, title, 40, True, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        text(s, 1.2, 4.2, 9.5, 0.7, subtitle, 20, False, "E0E7FF")
    text(s, 1.2, 6.5, 9, 0.5, f"{date.today():%B %d, %Y}   |   {BRAND}", 12, False, "C7D2FE")

    # --- Agenda
    if sections:
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, 13.333, 1.1, PALETTE[1])
        text(s, 0.6, 0.15, 12, 0.8, "Agenda", 30, True, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        items = [x["title"] for x in sections][:12]
        for i, name in enumerate(items):
            col, rowi = divmod(i, 6)
            x, y = 0.8 + col * 6.2, 1.6 + rowi * 0.85
            c = rect(s, x, y, 0.55, 0.55, PALETTE[i % len(PALETTE)], MSO_SHAPE.OVAL)
            c.text_frame.text = str(i + 1)
            para = c.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            para.runs[0].font.size, para.runs[0].font.bold = Pt(14), True
            para.runs[0].font.color.rgb = rgb("FFFFFF")
            text(s, x + 0.75, y, 5.2, 0.55, re.sub(r"^\d+[.)]\s*", "", name), 17, False, DARK,
                 anchor=MSO_ANCHOR.MIDDLE)

    def paginate(blocks):
        pages, cur, lines = [], [], 0

        def flush():
            nonlocal cur, lines
            if cur:
                pages.append(("text", cur))
            cur, lines = [], 0

        for b in blocks:
            if b["t"] == "table":
                hdr, body = b["rows"][0], b["rows"][1:]
                if cur and lines <= 6 and len(body) <= 6:
                    pages.append(("mixed", (cur, [hdr] + body)))
                    cur, lines = [], 0
                    continue
                flush()
                for i in range(0, max(len(body), 1), 8):
                    pages.append(("table", [hdr] + body[i:i + 8]))
            else:
                width = 90 - (10 if b.get("level") else 0)
                n = max(1, math.ceil(len(plain(b["text"])) / width)) + (1 if b["t"] == "h" else 0)
                if lines + n > 13:
                    flush()
                cur.append(b)
                lines += n
        flush()
        return pages

    def content_slide(name, accent, cont):
        sl = prs.slides.add_slide(blank)
        rect(sl, 0, 0, 13.333, 1.1, accent)
        rect(sl, 0, 1.1, 13.333, 0.06, PALETTE[3])
        label = re.sub(r"^\d+[.)]\s*", "", name) + (" (cont.)" if cont else "")
        text(sl, 0.6, 0.12, 12.2, 0.86, label, 28, True, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        rect(sl, 0, 7.18, 13.333, 0.32, LIGHT_ROW)
        text(sl, 0.6, 7.18, 8, 0.32, f"{BRAND}  |  {DISCLAIMER}", 9, False, "6B7280", anchor=MSO_ANCHOR.MIDDLE)
        return sl

    def draw_text(sl, blocks, accent):
        tb = sl.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.5))
        tf = tb.text_frame
        tf.word_wrap = True
        first, used = True, 0
        for b in blocks:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(8)
            used += max(1, math.ceil(len(plain(b["text"])) / 92)) + (1 if b["t"] == "h" else 0)
            if b["t"] == "h":
                rich(p, b["text"], 20, accent)
                for r in p.runs:
                    r.font.bold = True
            elif b["t"] in ("bullet", "num"):
                mark = "•  " if b["t"] == "bullet" else f"{b['n']}.  "
                if b.get("level"):
                    mark = "        –  "
                r = p.add_run()
                r.text = mark
                r.font.size, r.font.bold = Pt(18), True
                r.font.color.rgb = rgb(accent)
                rich(p, b["text"], 18)
            else:
                rich(p, b["text"], 18)
        return used

    def draw_table(sl, rows, accent, top):
        ncols = len(rows[0])
        size = 15 if ncols <= 3 else 13 if ncols <= 4 else 12 if ncols <= 6 else 10
        gt = sl.shapes.add_table(len(rows), ncols, Inches(0.6), Inches(top), Inches(12.1),
                                 Inches(0.55 * len(rows)))
        tbl = gt.table
        for ri, rowvals in enumerate(rows):
            for ci, v in enumerate(rowvals):
                cell = tbl.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(accent if ri == 0 else ("FFFFFF" if ri % 2 else LIGHT_ROW))
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                cell.text_frame.word_wrap = True
                p = cell.text_frame.paragraphs[0]
                rich(p, v, size, "FFFFFF" if ri == 0 else DARK)
                if ri == 0:
                    for r in p.runs:
                        r.font.bold = True

    for si, sec in enumerate(sections):
        accent = PALETTE[si % len(PALETTE)]
        pages = paginate(sec["blocks"]) or [("text", [])]
        for pi, (kind, data) in enumerate(pages):
            sl = content_slide(sec["title"], accent, pi > 0)
            if kind == "mixed":
                blocks_, rows_ = data
                used = draw_text(sl, blocks_, accent)
                draw_table(sl, rows_, accent, 1.55 + 0.42 * used + 0.2)
            elif kind == "text":
                draw_text(sl, data, accent)
            else:
                draw_table(sl, data, accent, 1.6)

    # --- Closing slide
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, PALETTE[1])
    rect(s, 9.5, -1.5, 5.5, 5.5, PALETTE[2], MSO_SHAPE.OVAL)
    rect(s, -1.5, 4.8, 4.5, 4.5, PALETTE[0], MSO_SHAPE.OVAL)
    text(s, 1.2, 2.5, 10, 1.2, "Thank you", 54, True, "FFFFFF")
    text(s, 1.2, 3.9, 10, 1.0, DISCLAIMER, 16, False, "E0E7FF")
    text(s, 1.2, 6.5, 9, 0.5, BRAND, 12, False, "C7D2FE")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
_PDF_MAP = {"→": "->", "←": "<-", "≥": ">=", "≤": "<=", "≈": "~", "✓": "+", "✔": "+", "–": "-", "—": "-",
            "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...", " ": " "}


def _pdf_text(t):
    t = clean(t)
    for k, v in _PDF_MAP.items():
        t = t.replace(k, v)
    t = t.encode("cp1252", "ignore").decode("cp1252")
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)


def to_pdf(md, title, subtitle):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    sections, h1 = split_sections(parse_markdown(md))
    title = title or h1 or "Market Research Report"
    hexc = lambda h: colors.HexColor("#" + h)

    margin = 0.75 * inch
    page_w, _ = A4
    avail = page_w - 2 * margin

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14, textColor=hexc(DARK),
                          spaceAfter=5)
    cell_style = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11, spaceAfter=0)
    head_cell = ParagraphStyle("hcell", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(hexc(PALETTE[0]))
        canvas.rect(0, A4[1] - 8, A4[0] * 0.4, 8, stroke=0, fill=1)
        canvas.setFillColor(hexc(PALETTE[2]))
        canvas.rect(A4[0] * 0.4, A4[1] - 8, A4[0] * 0.3, 8, stroke=0, fill=1)
        canvas.setFillColor(hexc(PALETTE[3]))
        canvas.rect(A4[0] * 0.7, A4[1] - 8, A4[0] * 0.3, 8, stroke=0, fill=1)
        canvas.setStrokeColor(hexc("D9DCF5"))
        canvas.line(margin, 0.6 * inch, A4[0] - margin, 0.6 * inch)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(hexc("6B7280"))
        canvas.drawString(margin, 0.42 * inch, f"{BRAND}  |  {DISCLAIMER}")
        canvas.drawRightString(A4[0] - margin, 0.42 * inch, f"Page {doc.page}")
        canvas.restoreState()

    story = []
    # cover banner
    t_style = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=26, leading=31, textColor=colors.white)
    s_style = ParagraphStyle("s", fontName="Helvetica", fontSize=12.5, leading=16, textColor=hexc("E0E7FF"))
    d_style = ParagraphStyle("d", fontName="Helvetica", fontSize=9.5, leading=12, textColor=hexc("C7D2FE"))
    cover = Table([[Paragraph(_pdf_text(title), t_style)],
                   [Paragraph(_pdf_text(subtitle or ""), s_style)],
                   [Paragraph(f"Prepared {date.today():%B %d, %Y}  |  {BRAND}", d_style)]], colWidths=[avail])
    cover.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), hexc(PALETTE[0])),
                               ("LEFTPADDING", (0, 0), (-1, -1), 18), ("RIGHTPADDING", (0, 0), (-1, -1), 18),
                               ("TOPPADDING", (0, 0), (0, 0), 22), ("BOTTOMPADDING", (0, -1), (-1, -1), 20),
                               ("TOPPADDING", (0, 1), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -2), 4)]))
    story.append(cover)
    strip = Table([["", "", "", ""]], colWidths=[avail / 4] * 4, rowHeights=[7])
    strip.setStyle(TableStyle([("BACKGROUND", (i, 0), (i, 0), hexc(PALETTE[(i + 1) % len(PALETTE)])) for i in range(4)]))
    story += [strip, Spacer(1, 14)]

    for si, sec in enumerate(sections):
        accent = PALETTE[si % len(PALETTE)]
        h2 = ParagraphStyle(f"h2{si}", fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=hexc(accent),
                            spaceBefore=12, spaceAfter=2)
        h3 = ParagraphStyle(f"h3{si}", fontName="Helvetica-Bold", fontSize=11.5, leading=15, textColor=hexc(accent),
                            spaceBefore=7, spaceAfter=3)
        story.append(KeepTogether([Paragraph(_pdf_text(sec["title"]), h2),
                                   HRFlowable(width="100%", thickness=1.6, color=hexc(accent), spaceAfter=6)]))
        for b in sec["blocks"]:
            t = b["t"]
            if t == "h":
                story.append(Paragraph(_pdf_text(b["text"]), h3))
            elif t == "p":
                story.append(Paragraph(_pdf_text(b["text"]), body))
            elif t in ("bullet", "num"):
                lvl = b.get("level", 0)
                st = ParagraphStyle("li", parent=body, leftIndent=16 + 14 * lvl, bulletIndent=4 + 14 * lvl,
                                    spaceAfter=3, bulletColor=hexc(accent), bulletFontName="Helvetica-Bold")
                story.append(Paragraph(_pdf_text(b["text"]), st,
                                       bulletText="•" if t == "bullet" else f"{b['n']}."))
            elif t == "table":
                rows = b["rows"]
                data = [[Paragraph(_pdf_text(c), head_cell if ri == 0 else cell_style) for c in r]
                        for ri, r in enumerate(rows)]
                tbl = Table(data, colWidths=[avail / len(rows[0])] * len(rows[0]), repeatRows=1)
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), hexc(accent)),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hexc(LIGHT_ROW)]),
                    ("GRID", (0, 0), (-1, -1), 0.4, hexc("D9DCF5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
                story += [Spacer(1, 3), tbl, Spacer(1, 8)]

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=0.7 * inch,
                      bottomMargin=0.85 * inch, title=title, author=BRAND).build(
        story, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


# ------------------------------------------------------------------ public API
def build_all(md, title, subtitle):
    """Returns ({'docx': bytes, ...}, {'docx': 'error text', ...})."""
    makers = {"docx": to_docx, "xlsx": to_xlsx, "pptx": to_pptx, "pdf": to_pdf}
    files, errors = {}, {}
    for key, fn in makers.items():
        try:
            files[key] = fn(md, title, subtitle)
        except Exception as e:  # keep the other formats working
            errors[key] = f"{type(e).__name__}: {e}"
    return files, errors
